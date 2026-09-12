"""미국 금융 스트레스 파이프라인의 외부 공식자료 수집·파싱 어댑터.

이 모듈은 자료를 월·주 단위 값으로 정규화하는 일까지만 담당한다.
지수 산식과 Supabase 저장 순서는 ``financial_stress_pipeline``에 둔다.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta

import openpyxl
import requests

from common import fetch_fred_observations


TIMEOUT_SECONDS = 45
OFFICIAL_DATA_HEADERS = {"User-Agent": "MacroWatch/1.0 (+https://hoorash4.github.io/macrowatch/)"}
EBP_CSV_URL = "https://www.federalreserve.gov/econres/notes/feds-notes/ebp_csv.csv"
CMDI_XLSX_URL = "https://www.newyorkfed.org/medialibrary/research/interactives/cmdi/downloads/Market%20CMDI.xlsx"


def month_start(value: date) -> str:
    return value.replace(day=1).isoformat()


def _fred_observations(series_id: str, api_key: str, start: date, end: date):
    return fetch_fred_observations(
        series_id,
        api_key,
        start=start.isoformat(),
        end=end.isoformat(),
        timeout=TIMEOUT_SECONDS,
    )


def fetch_fred_month_end(series_id: str, api_key: str, start: date, end: date) -> dict[str, float]:
    """각 달의 마지막 유효 일별 관측값을 반환한다."""
    values: dict[str, tuple[str, float]] = {}
    for observation in _fred_observations(series_id, api_key, start, end):
        raw_value, observed_on = observation.get("value"), observation.get("date")
        if raw_value in (None, ".") or not isinstance(observed_on, str):
            continue
        try:
            value = float(raw_value)
        except (TypeError, ValueError):
            continue
        month = observed_on[:7] + "-01"
        if month not in values or observed_on > values[month][0]:
            values[month] = (observed_on, value)
    return {month: value for month, (_observed_on, value) in values.items()}


def fetch_fred_week_end(series_id: str, api_key: str, start: date, end: date) -> dict[str, float]:
    """금요일 종료 주간별 마지막 유효 관측값을 반환한다."""
    values: dict[str, tuple[str, float]] = {}
    for observation in _fred_observations(series_id, api_key, start, end):
        raw_value, observed_on = observation.get("value"), observation.get("date")
        if raw_value in (None, ".") or not isinstance(observed_on, str):
            continue
        try:
            observed_date, value = date.fromisoformat(observed_on), float(raw_value)
        except (TypeError, ValueError):
            continue
        week = (observed_date + timedelta(days=4 - observed_date.weekday())).isoformat()
        if week not in values or observed_on > values[week][0]:
            values[week] = (observed_on, value)
    return {week: value for week, (_observed_on, value) in values.items()}


def fetch_ebp_monthly(start: date, end: date) -> dict[str, float]:
    """연준의 월별 Excess Bond Premium CSV를 읽는다."""
    response = requests.get(EBP_CSV_URL, headers=OFFICIAL_DATA_HEADERS, timeout=TIMEOUT_SECONDS)
    response.raise_for_status()
    reader = csv.DictReader(io.StringIO(response.content.decode("utf-8-sig", errors="replace")))
    values: dict[str, float] = {}
    for row in reader:
        normalized = {str(key).strip().lower(): value for key, value in row.items() if key}
        raw_date = normalized.get("date") or normalized.get("observation_date")
        if raw_date is None:
            raw_date = next((value for key, value in normalized.items() if "date" in key), None)
        raw_value = normalized.get("ebp")
        if raw_value is None:
            raw_value = next(
                (value for key, value in normalized.items() if key.endswith("ebp") or "excess bond premium" in key),
                None,
            )
        observed_on = _parse_date(raw_date)
        try:
            value = float(str(raw_value))
        except (TypeError, ValueError):
            continue
        if observed_on is not None and start <= observed_on <= end:
            values[month_start(observed_on)] = value
    if not values:
        raise RuntimeError("연준 EBP CSV에서 월별 값을 찾지 못했습니다.")
    return values


def _parse_date(raw_value: object) -> date | None:
    if isinstance(raw_value, datetime):
        return raw_value.date()
    if isinstance(raw_value, date):
        return raw_value
    raw_text = str(raw_value).split(" ")[0]
    for pattern in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw_text, pattern).date()
        except ValueError:
            continue
    return None


def fetch_cmdi_monthly(start: date, end: date) -> dict[str, float]:
    """뉴욕연은 Market CMDI 워크북을 메모리에서 읽는다."""
    response = requests.get(CMDI_XLSX_URL, headers=OFFICIAL_DATA_HEADERS, timeout=TIMEOUT_SECONDS)
    response.raise_for_status()
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    values: dict[str, tuple[date, float]] = {}
    try:
        for sheet in workbook.worksheets:
            header_row = date_column = value_column = None
            for row_number, header in enumerate(sheet.iter_rows(min_row=1, max_row=100, values_only=True), start=1):
                labels = [str(value or "").strip().lower() for value in header]
                try:
                    date_column = next(
                        index for index, label in enumerate(labels)
                        if label == "date" or "date" in label or label in {"eow_friday", "week_end", "week ending"}
                    )
                    value_column = next(
                        index for index, label in enumerate(labels)
                        if ("market" in label and "cmdi" in label) or label == "market"
                    )
                    header_row = row_number
                    break
                except StopIteration:
                    continue
            if header_row is None or date_column is None or value_column is None:
                continue
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if len(row) <= max(date_column, value_column):
                    continue
                observed_on = _parse_date(row[date_column])
                try:
                    value = float(row[value_column])
                except (TypeError, ValueError):
                    continue
                if observed_on is not None and start <= observed_on <= end:
                    month = month_start(observed_on)
                    if month not in values or observed_on > values[month][0]:
                        values[month] = (observed_on, value)
    finally:
        workbook.close()
    if not values:
        raise RuntimeError("뉴욕연은 CMDI XLSX에서 Market CMDI 값을 찾지 못했습니다.")
    return {month: value for month, (_observed_on, value) in values.items()}
