"""Official Supreme Court monthly corporate-rehabilitation statistics source."""
from __future__ import annotations

from datetime import date
from io import BytesIO
import json
from typing import Iterator

import requests

from common import request_with_retry

BASE_URL = "https://portal.scourt.go.kr"
USER_AGENT = "Mozilla/5.0 (compatible; MacroWatch/1.0; +https://hoorash4.github.io/macrowatch/)"
LARGE_CATEGORY = "G01"  # 민사
MIDDLE_CATEGORY = "T09"  # 도산관련
SMALL_CATEGORY = "S05"  # 도산>회생합의사건
SOURCE = "SupremeCourt:PGP441M01/S05"


def _post(path: str, body: dict) -> dict:
    response = request_with_retry(lambda: requests.post(
        BASE_URL + path,
        json=body,
        headers={"User-Agent": USER_AGENT, "Content-Type": "application/json"},
        timeout=45,
    ))
    response.raise_for_status()
    payload = response.json()
    if payload.get("status") != 200:
        raise RuntimeError(f"Supreme Court API returned status {payload.get('status')}")
    return payload.get("data") or {}


def _available_months(year: int) -> list[int]:
    data = _post("/pgp/pgp441/selectCortStatsMonth.on", {
        "dma_year": {
            "targetYear": str(year),
            "cortStatsItmLclId": LARGE_CATEGORY,
            "cortStatsItmMclId": MIDDLE_CATEGORY,
        }
    })
    months: list[int] = []
    for row in data.get("dlt_cortStatsMonth") or []:
        try:
            month = int(row.get("aojStatsMm"))
        except (TypeError, ValueError):
            continue
        if 1 <= month <= 12:
            months.append(month)
    return sorted(set(months))


def _workbook_metadata(year: int, month: int) -> dict:
    data = _post("/pgp/pgp441/selectCortStatsMmDtl.on", {
        "dma_search": {
            "cortStatsItmLclId": LARGE_CATEGORY,
            "cortStatsItmMclId": MIDDLE_CATEGORY,
            "cortStatsItmSclId": SMALL_CATEGORY,
            "aojStatsYr": str(year),
            "aojStatsMm": f"{month:02d}",
            "ojdpAtflDvsCd": "07",
        },
        "dma_pageInfo": {"pageNo": 1, "pageSize": 10, "totalYn": "Y", "totalCnt": 10},
    })
    rows = data.get("dlt_aojAlmnLst") or []
    for row in rows:
        if str(row.get("cortStatsItmSclNm") or "").endswith("회생합의사건"):
            if row.get("excelNm") and row.get("excelPath"):
                return row
    raise RuntimeError(f"Supreme Court rehabilitation workbook missing for {year:04d}-{month:02d}")


def _download_workbook(meta: dict) -> bytes:
    response = request_with_retry(lambda: requests.get(
        BASE_URL + "/pgp/pgp003/downloadEml.on",
        params={
            "fileNm": meta["excelNm"],
            "filePathNm": meta["excelPath"],
            "saveFileName": meta.get("excelNtatcAtflNm") or meta["excelNm"],
        },
        headers={"User-Agent": USER_AGENT},
        timeout=45,
    ))
    response.raise_for_status()
    if not response.content.startswith(b"PK"):
        raise RuntimeError("Supreme Court rehabilitation download is not an XLSX workbook")
    return response.content


def _load_workbook(content: bytes):
    """Import the optional XLSX reader only at the source boundary.

    Some unrelated repository tests install lightweight module doubles while the full suite is
    importing. Keeping openpyxl out of module import time prevents those doubles from affecting
    source discovery; production parsing still uses the pinned openpyxl dependency.
    """
    from openpyxl import load_workbook

    return load_workbook(BytesIO(content), data_only=True, read_only=True)


def _parse_monthly_filings(content: bytes) -> int:
    """Return the nationwide monthly filing count from the official 회생합의 workbook."""
    workbook = _load_workbook(content)
    try:
        sheet = workbook["회생합의"] if "회생합의" in workbook.sheetnames else workbook.worksheets[0]
        for values in sheet.iter_rows(values_only=True):
            if not values:
                continue
            label = str(values[0] or "").replace(" ", "").strip()
            if label != "총계":
                continue
            value = values[1] if len(values) > 1 else None
            if isinstance(value, bool):
                break
            try:
                count = int(float(value))
            except (TypeError, ValueError):
                break
            if count < 0:
                break
            return count
    finally:
        workbook.close()
    raise RuntimeError("Supreme Court rehabilitation workbook has no valid 총계/접수 value")


def _row(year: int, month: int, value: int) -> dict:
    return {
        "series_code": "KR_CORP_REHAB",
        "observation_date": f"{year:04d}-{month:02d}-01",
        "value": value,
        "frequency": "M",
        "source": SOURCE,
    }


def fetch_korea_corporate_rehab_month(year: int, month: int) -> dict:
    meta = _workbook_metadata(year, month)
    count = _parse_monthly_filings(_download_workbook(meta))
    return _row(year, month, count)


def iter_korea_corporate_rehab_rows(start: date, end: date) -> Iterator[dict]:
    """Yield all available official monthly rows, isolating individual source failures."""
    first = date(start.year, start.month, 1)
    last = date(end.year, end.month, 1)
    attempted = 0
    succeeded = 0
    errors: list[str] = []
    for year in range(first.year, last.year + 1):
        try:
            months = _available_months(year)
        except Exception as error:
            errors.append(f"{year}: {error.__class__.__name__}: {error}")
            continue
        for month in months:
            observed = date(year, month, 1)
            if observed < first or observed > last:
                continue
            attempted += 1
            try:
                row = fetch_korea_corporate_rehab_month(year, month)
            except Exception as error:
                errors.append(f"{year:04d}-{month:02d}: {error.__class__.__name__}: {error}")
                continue
            succeeded += 1
            yield row
    if errors:
        print(json.dumps({
            "stage": "court_rehabilitation_source_partial_errors",
            "attempted": attempted,
            "succeeded": succeeded,
            "errors": errors,
        }, ensure_ascii=False))
    if attempted and not succeeded:
        raise RuntimeError("Supreme Court rehabilitation source returned no usable monthly rows")


def fetch_korea_corporate_rehab_rows(start: date, end: date) -> list[dict]:
    return list(iter_korea_corporate_rehab_rows(start, end))
